"""筛选页测试：开票年月过滤、选择计数与 Excel 导出。

年月解析口径：
  · 各种日期写法（标准/中文/斜杠）→ "YYYY-MM"
  · 旧合同数据无 invoice_date 时回落 sign_date
  · 不可解析 → 空串（不参与下拉选项）

运行：.venv/Scripts/python.exe -m pytest tests/test_filter_page.py -v
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from ui.pages.filter_page import COL_ACTION, COL_CHECK, COL_NAME, record_month, year_month


def test_year_month_parses_common_formats():
    assert year_month("2026-09-09") == "2026-09"
    assert year_month("2026-9-9") == "2026-09"
    assert year_month("2026年09月09日") == "2026-09"
    assert year_month("2026年9月9日") == "2026-09"
    assert year_month("2026/09/09") == "2026-09"
    assert year_month("2026.09.09") == "2026-09"


def test_year_month_rejects_unparsable():
    assert year_month(None) == ""
    assert year_month("") == ""
    assert year_month("无日期") == ""
    assert year_month("HT20260901") == ""  # 合同编号这类不该被当成日期


def test_record_month_prefers_invoice_date_then_sign_date():
    assert record_month({"invoice_date": "2026-09-09"}) == "2026-09"
    assert record_month({"invoice_date": None, "sign_date": "2025-12-31"}) == "2025-12"
    assert record_month({"contract_no": "HT20260901"}) == ""


# --------------------------------------------------------------- 页面与导出（UI 流程）


def _select_all(page, checked: bool = True) -> None:
    for row in range(page._table.rowCount()):
        page._table.cellWidget(row, COL_CHECK).setChecked(checked)


def _row_of(page, file_name: str) -> int:
    """按文件名定位行（列表按导入时间倒序，不能假定固定行号）。"""
    for row in range(page._table.rowCount()):
        if page._table.item(row, COL_NAME).text() == file_name:
            return row
    raise AssertionError(f"列表中找不到 {file_name}")


# 演示数据：可疑 2 条（manual_review / warning）、准确 1、失败 1
CONFIRM_DOCS = (
    ("review.pdf", "manual_review"),
    ("warning.pdf", "warning"),
    ("success.pdf", "success"),
    ("failed.pdf", "failed"),
)


@pytest.fixture
def qt_app():
    """离屏 QApplication（platform 由 conftest 统一设为 offscreen）。"""
    from PySide6.QtWidgets import QApplication

    return QApplication.instance() or QApplication([])


def _make_page(db_path, docs=()) -> object:
    """建库（可选写入演示记录）并构造筛选页。"""
    from database.db import get_engine, init_db, make_session_factory
    from models.document import Document
    from ui.pages.filter_page import FilterPage

    engine = get_engine(db_path)
    init_db(engine)
    if docs:
        with make_session_factory(engine)() as session:
            for name, status in docs:
                session.add(
                    Document(
                        file_name=name,
                        file_path=f"C:/{name}",
                        file_hash=f"hash-{name}",
                        status=status,
                        error_reason=None if status == "success" else "需人工核对",
                    )
                )
            session.commit()
    return FilterPage(str(db_path))


@pytest.fixture
def confirm_page(qt_app, tmp_path):
    """带 4 条记录的筛选页（见 CONFIRM_DOCS）。"""
    db_path = tmp_path / "app.db"
    return _make_page(db_path, CONFIRM_DOCS), db_path


@pytest.fixture
def empty_page(qt_app, tmp_path):
    """空库筛选页：用于验证空列表下的文案/按钮表现。"""
    db_path = tmp_path / "empty.db"
    return _make_page(db_path), db_path


def test_one_click_confirm_button_removed(confirm_page):
    page, _ = confirm_page
    from PySide6.QtWidgets import QPushButton

    assert not hasattr(page, "_confirm_btn")
    assert all("一键确认" not in button.text() for button in page.findChildren(QPushButton))


def test_action_column_is_part_of_main_table(confirm_page):
    page, _ = confirm_page

    assert not hasattr(page, "_action_table")
    assert page._table.horizontalHeaderItem(COL_ACTION).text() == "操作"
    assert page._table.cellWidget(0, COL_ACTION) is not None


def test_double_click_any_cell_opens_row_detail(confirm_page):
    page, _ = confirm_page
    row = _row_of(page, "review.pdf")
    emitted = []
    page.detail_requested.connect(lambda doc_id, name: emitted.append((doc_id, name)))

    page._table.cellDoubleClicked.emit(row, COL_NAME)

    assert emitted and emitted[0][1] == "review.pdf"


def test_row_context_menu_contains_actions(confirm_page):
    page, _ = confirm_page

    texts = [action.text() for action in page._row_menu(0).actions() if not action.isSeparator()]

    assert texts == ["查看详情", "复制整行", "删除"]


def test_export_includes_success_and_suspect_but_skips_failed(
    confirm_page, monkeypatch, tmp_path
):
    page, _ = confirm_page
    output = tmp_path / "selected.xlsx"
    from PySide6.QtWidgets import QFileDialog, QMessageBox

    prompts = []
    monkeypatch.setattr(
        QMessageBox,
        "exec",
        lambda dialog: (
            prompts.append((dialog.text(), dialog.minimumWidth()))
            or QMessageBox.StandardButton.Yes
        ),
    )

    monkeypatch.setattr(
        QFileDialog,
        "getSaveFileName",
        staticmethod(lambda *args, **kwargs: (str(output), "Excel 工作簿 (*.xlsx)")),
    )
    _select_all(page)

    page._export_selected()

    assert "准确数据：1 条" in prompts[0][0]
    assert "可疑数据：2 条" in prompts[0][0]
    assert "不可导出并跳过：1 条" in prompts[0][0]
    assert prompts[0][1] == 300
    from openpyxl import load_workbook

    ws = load_workbook(output, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "文件名"
    assert rows[0][-1] == "状态"
    assert {row[0]: row[-1] for row in rows[1:]} == {
        "review.pdf": "可疑/待校验",
        "warning.pdf": "可疑/待校验",
        "success.pdf": "准确",
    }
    assert "跳过 1 条" in page._toast.text()


def test_export_cancel_stops_before_file_dialog(confirm_page, monkeypatch):
    page, _ = confirm_page
    from PySide6.QtWidgets import QFileDialog, QMessageBox

    save_dialog_opened = []
    monkeypatch.setattr(
        QMessageBox,
        "exec",
        lambda dialog: QMessageBox.StandardButton.No,
    )
    monkeypatch.setattr(
        QFileDialog,
        "getSaveFileName",
        staticmethod(lambda *args, **kwargs: save_dialog_opened.append(True) or ("", "")),
    )
    _select_all(page)

    page._export_selected()

    assert save_dialog_opened == []


# ------------------------------------------------------- "已选择 N 项" 计数（UI 流程）


def _count_text(page) -> str:
    """去掉 <b> 标签后的计数文案，例如 "已选择 3 项"。"""
    return re.sub(r"<[^>]+>", "", page._select_count_label.text())


def test_select_count_starts_at_zero(confirm_page):
    page, _ = confirm_page
    assert _count_text(page) == "已选择 0 项"
    assert page._select_count_label.isHidden() is False  # 有数据时显示


def test_select_count_updates_on_row_check(confirm_page):
    """手动勾选/取消勾选要实时刷新（修复前只在 reload 后刷新，会一直停在 0）。"""
    page, _ = confirm_page
    check = page._table.cellWidget(_row_of(page, "warning.pdf"), COL_CHECK)

    check.setChecked(True)
    assert _count_text(page) == "已选择 1 项"

    page._table.cellWidget(_row_of(page, "review.pdf"), COL_CHECK).setChecked(True)
    assert _count_text(page) == "已选择 2 项"

    check.setChecked(False)
    assert _count_text(page) == "已选择 1 项"


def test_select_count_updates_on_select_all(confirm_page):
    page, _ = confirm_page

    page._select_all_btn.setChecked(True)
    assert _count_text(page) == "已选择 4 项"

    page._select_all_btn.setChecked(False)
    assert _count_text(page) == "已选择 0 项"


def test_select_count_refreshes_after_data_reload(confirm_page):
    """刷新数据后勾选态清空、计数归零（状态筛选切换走到 reload）。"""
    page, _ = confirm_page
    page._select_all_btn.setChecked(True)
    assert _count_text(page) == "已选择 4 项"

    page._status_filter.setCurrentIndex(page._status_filter.findData("review"))
    page.reload()

    assert page._table.rowCount() == 2
    assert _count_text(page) == "已选择 0 项"


def test_select_count_label_hidden_when_list_empty(empty_page):
    """空列表不显示"已选择 0 项"（无意义文案）。"""
    page, _ = empty_page
    assert page._table.rowCount() == 0
    assert page._select_count_label.isHidden() is True

